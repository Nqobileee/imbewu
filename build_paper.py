# -*- coding: utf-8 -*-
"""Builds the research paper (DOCX + PDF) in JRISS format from one shared content model."""
import os
from PIL import Image as PILImage

TITLE = ("Offline Artificial Intelligence for Precision Crop Disease Diagnosis "
         "in Low-Connectivity Rural Environments")
DOCX_OUT = "Offline_AI_Crop_Disease_Diagnosis.docx"
PDF_OUT  = "Offline_AI_Crop_Disease_Diagnosis.pdf"

# ----------------------------------------------------------------------------- content model
# block tuples: (kind, *payload)
B = []
def hdr(lines): B.append(("hdr", lines))
def title(t):   B.append(("title", t))
def auth(lines):B.append(("auth", lines))
def meta(lines):B.append(("meta", lines))
def ab(t):      B.append(("abstract", t))
def kw(t):      B.append(("keywords", t))
def h1(t):      B.append(("h1", t))
def h2(t):      B.append(("h2", t))
def p(t):       B.append(("body", t))
def eq(t, n):   B.append(("eq", t, n))
def tbl(cap, head, rows): B.append(("table", cap, head, rows))
def fig(path, cap): B.append(("figure", path, cap))
def refs(items): B.append(("refs", items))

hdr(["Journal of Research and Innovation for Sustainable Society (JRISS)",
     "Volume 6, Issue 1, 2026", "ISSN: 2668-0416"])
title(TITLE)
auth(["Andile Prince Bhebhe¹ and Edith Nqobile Muyambiri²",
      "¹Liquid Intelligent Technologies, Bulawayo, Zimbabwe; Omnishi AI; and National University of Science and Technology alumni",
      "²Harare Institute of Technology, Harare, Zimbabwe; and ZB Innovation Hub, Harare, Zimbabwe",
      "E-mail: andileprincebhebhe@gmail.com ; nqobilemuyambiri@gmail.com"])
meta(["Theme: Area 8, Empowering SMEs with 4IR and AI for Inclusive and Sustainable Growth",
      "Category: Prototype Demonstration"])

ab(("Crop losses to pest and disease outbreaks threaten food security across rural Zimbabwe, "
 "where agriculture supports roughly 70% of the population and where Fourth Industrial Revolution "
 "technologies are increasingly promoted as drivers of inclusive development and progress towards "
 "the Sustainable Development Goals. Existing digital advisory tools predominantly "
 "require internet connectivity and are tailored to large commercial operations, leaving "
 "smallholder farmers in low-connectivity areas dependent on overstretched extension officers and "
 "informal, unverified advice. This "
 "paper presents an offline, on-device artificial-intelligence pipeline for precision crop disease "
 "diagnosis, demonstrated on Northern Corn Leaf Blight (NCLB) in maize. A locally trained "
 "convolutional neural network performs inference directly on a smartphone, so a farmer can "
 "photograph an affected leaf and obtain a diagnosis with no network access. The diagnostic engine "
 "is built by transfer learning on a lightweight MobileNetV3-Small backbone pretrained on ImageNet "
 "and is evaluated on a dataset of 1,821 field-collected "
 "maize-leaf images in which the non-NCLB class consists of other maize diseases rather than "
 "healthy foliage, forcing the model to distinguish NCLB's characteristic tan, cigar-shaped lesions "
 "from other disease symptoms. Two routes are compared on an identical, leakage-free, group-aware "
 "test partition: classical classifiers trained on frozen 576-dimensional embeddings, and an "
 "end-to-end fine-tuned network with data augmentation and class-weighted loss. The fine-tuned "
 "model attains test accuracy 96.0%, balanced accuracy 0.959, macro-F1 0.960 and ROC-AUC 0.989, "
 "while a support-vector classifier on frozen embeddings reaches test ROC-AUC 0.998. Grad-CAM "
 "attributions confirm that predictions are driven by lesion regions. "
 "Because the quantised model is only a few megabytes and runs on a CPU, the system is deployable "
 "on low-cost phones entirely offline, with treatment guidance delivered in English, Shona and "
 "Ndebele. The work aligns with Zimbabwe's National AI Strategy (2026-2030) and advances SDG 2 "
 "(Zero Hunger) and SDG 9 (Industry, Innovation and Infrastructure)."))
kw(("crop disease detection, Northern Corn Leaf Blight, transfer learning, MobileNetV3, "
    "on-device inference, offline AI, precision agriculture, convolutional neural networks, "
    "Grad-CAM, food security, Zimbabwe"))

# ---------------- 1 Background
h1("1 Background")
p("Zimbabwe was once regarded as a regional breadbasket, but that standing has eroded markedly, and "
  "with it the country's food security. Agriculture remains the backbone of the economy, employing "
  "roughly 60-70% of the workforce and supporting about 70% of the population, predominantly through "
  "smallholder maize (FAO, 2021a; Milling MEA, 2024). Yet repeated shocks have turned a former "
  "surplus producer into a recurrent importer: the 2023/24 El Nino drought cut national maize output "
  "to roughly 1.5 million tonnes, leaving a deficit of nearly 700,000 tonnes (USDA FAS, 2024) and "
  "prompting a national state of disaster, with around 6 million people projected to be food "
  "insecure during the 2024-2025 lean season (OCHA, 2024). This fragility sits within a much larger "
  "crop-health problem. Globally, the FAO estimates that up to 40% of crop production is lost to "
  "pests and diseases each year, at a cost exceeding US$220 billion (FAO, 2021b), and maize alone "
  "loses an estimated 22.5% of its potential yield to pathogens and pests, with the heaviest burden "
  "falling on food-deficit regions with fast-growing populations (Savary et al., 2019). In "
  "sub-Saharan Africa, where maize is a staple for some 300 million people and is grown "
  "overwhelmingly by smallholders (CABI, 2024), the fall armyworm alone could destroy 8.3-20.6 "
  "million tonnes of maize a year across twelve countries in the absence of control (Day et al., "
  "2017); in Zimbabwe it has been projected to cut national output by around 264,000 tonnes "
  "annually (Rwomushana et al., 2018), with affected households roughly 12% more likely to "
  "experience hunger (Tambo et al., 2020). Among foliar diseases of maize, Northern Corn Leaf Blight "
  "(NCLB), caused by Exserohilum turcicum, is one of the most damaging worldwide, capable of "
  "reducing grain yield by 30-50% when it establishes early in the season (Jackson-Ziems, 2014), "
  "which makes its rapid, accurate identification a high-value intervention.")
p("Closing this gap in a low-connectivity rural setting is constrained by the scarcity of "
  "agricultural extension officers, who cannot meet real-time demand, leaving farmers dependent on "
  "informal and frequently unverified advice (FAO, 2021a). This is precisely the need that national "
  "policy now prioritises. Zimbabwe's National AI Strategy (2026-2030) identifies agriculture as a "
  "priority sector and explicitly commits to leveraging AI to strengthen food security, build "
  "climate resilience, and support both smallholder and commercial farmers, including through crop "
  "disease prevention (Government of Zimbabwe, 2026; CIPESA, 2026); tellingly, the inaugural "
  "Zimbabwe AI Grand Challenge, the first edition of the strategy's flagship annual competition, is "
  "dedicated to food security (Government of Zimbabwe, 2026). An offline, on-device tool for early "
  "maize disease diagnosis therefore responds not only to a documented agronomic and humanitarian "
  "need, but to an explicitly stated national AI priority, advancing SDG 2 (Zero Hunger) and SDG 9 "
  "(Industry, Innovation and Infrastructure), and reflecting wider evidence that Fourth Industrial "
  "Revolution technologies can drive inclusive development in emerging economies (Mhlanga, 2021).")

# ---------------- 2 Introduction
h1("2 Introduction")
p("Building on the agronomic and policy context above, this paper addresses the technical problem "
  "of detecting NCLB from ordinary maize-leaf photographs. Early, accurate identification enables "
  "timely intervention, but the advisory capacity to deliver it at scale is lacking in "
  "low-connectivity rural areas, motivating an automated diagnostic tool that operates without a "
  "network connection and on the modest hardware farmers already own.")
p("Digital tools could in principle close this gap, but most existing solutions assume reliable "
  "internet connectivity and cloud inference, and are designed for large commercial operations. In "
  "rural areas with intermittent or absent network coverage these assumptions fail. An effective "
  "intervention must therefore operate offline, on the modest hardware farmers already own, while "
  "remaining accurate enough to be trusted. Advances in efficient convolutional neural networks "
  "(CNNs) and transfer learning make on-device image classification feasible: a compact model "
  "pretrained on a large corpus can be specialised to a narrow agricultural task with comparatively "
  "little labelled data (Zhuang et al., 2021).")
p("This study treats crop disease diagnosis as an image-classification problem and demonstrates an "
  "offline-first design on the binary task of distinguishing NCLB from other maize diseases. The "
  "choice of negatives is deliberate. A model trained to separate diseased from healthy leaves can "
  "succeed by detecting any departure from green, healthy tissue; such a model is of little clinical "
  "value because the difficult decisions in the field are between diseases that look superficially "
  "similar. By composing the non-NCLB class entirely from other maize diseases, the task forces the "
  "network to learn the specific morphology of blight lesions rather than a generic healthy-versus-"
  "diseased cue.")
p("The contribution of this paper is fourfold and aligns with the JRISS prototype-demonstration "
  "category. First, it curates a field-collected maize-leaf dataset in which the negative class "
  "comprises other diseases, providing a realistic and demanding benchmark. Second, it adopts a "
  "leakage-free, group-aware evaluation in which all images originating from the same physical leaf "
  "are confined to a single data partition, so reported scores reflect generalisation to unseen "
  "plants rather than memorisation. Third, it compares two transfer-learning routes, frozen-embedding "
  "classical classifiers and an end-to-end fine-tuned CNN, under identical evaluation. Fourth, it "
  "couples interpretability (Grad-CAM) with a concrete offline deployment architecture suitable for "
  "low-cost smartphones, in direct response to the policy direction of Zimbabwe's National AI "
  "Strategy (2026-2030).")
fig("outputs/figures/01_composition.png",
    "Figure 1. Dataset composition: class balance (NCLB versus other maize diseases) and the fixed "
    "train/validation/test partition. Source: study exploratory analysis, June 2026.")
fig("outputs/figures/samples_paper.png",
    "Figure 2. Representative dataset samples: NCLB lesions (top row) versus other maize diseases "
    "(bottom row), each centre-cropped and resized to 224x224 pixels. Source: study dataset.")

# ---------------- 2 Literature review
h1("3 Literature review")
p("Image-based plant disease classification has matured rapidly, and recent reviews synthesise the "
  "state of the art. Ngugi et al. (2021) survey image-processing and deep-learning techniques for "
  "automated leaf pest and disease recognition, and Shoaib et al. (2023) review advanced "
  "deep-learning models for plant disease detection; both report that convolutional neural networks "
  "now reach very high benchmark accuracy yet degrade markedly when tested on imagery captured under "
  "conditions different from training, the generalisation gap that remains the field's central "
  "challenge and that laboratory-style images with uniform backgrounds tend to mask.")
p("Transfer learning is the dominant strategy when labelled agricultural data are limited (Zhuang "
  "et al., 2021): a network pretrained on a large corpus is re-used as a feature extractor or "
  "initialisation for a smaller target task. Modern convolutional architectures and their "
  "accuracy-efficiency trade-offs are surveyed by Alzubaidi et al. (2021), while efficient families "
  "such as EfficientNetV2 (Tan and Le, 2021) and the MobileNet line are engineered for inference "
  "under tight compute and memory budgets, making compact backbones natural candidates for offline "
  "deployment. Recent plant-disease reviews confirm that such lightweight networks retain most of "
  "the accuracy of larger ones at a fraction of the cost (Shoaib et al., 2023).")
p("Within Africa, recent work has applied deep learning to staple-crop disease using increasingly "
  "field-realistic data. Sambasivam and Opiyo (2021) trained convolutional networks for cassava "
  "disease classification on an imbalanced East African dataset, directly tackling the "
  "class-imbalance that characterises field collections. Mensah et al. (2023) released the CCMT "
  "crop pest and disease dataset spanning maize, cassava, cashew and tomato imaged in Ghanaian "
  "fields, expanding the still-scarce pool of African training data. Waldamichael et al. (2022) "
  "reviewed machine learning for cereal-crop disease detection, including maize, and emphasised "
  "both the recurring domain gap between curated and in-field imagery and the value of lightweight, "
  "deployable models, two concerns central to the present design.")
p("For Zimbabwe and the wider Southern African region, image-based crop-disease diagnosis tailored "
  "to smallholder, low-connectivity conditions remains nascent, and offline, on-device tools for "
  "maize, the country's principal staple, are particularly scarce. A recent assessment of "
  "Zimbabwe's AI readiness by Dube et al. (2025), at the National University of Science and "
  "Technology in Bulawayo, found growing awareness and emerging applications in agriculture and "
  "banking alongside university innovation hubs, but persistent constraints in infrastructure, "
  "research funding, AI skills, and rural digital literacy, precisely the conditions an offline, "
  "low-cost tool is designed to navigate. This study addresses that gap and responds directly to "
  "Zimbabwe's National AI Strategy (2026-2030), which prioritises inclusive, locally relevant and "
  "explainable artificial intelligence for public-good applications, and to the national "
  "agriculture and food-systems transformation agenda (Government of Zimbabwe, 2026).")
tbl("Table 1. Definitions of key terms.",
    ["Key term", "Definition"],
    [["Northern Corn Leaf Blight (NCLB)",
      "A fungal foliar disease of maize (Exserohilum turcicum) producing long, tan to grey, "
      "cigar-shaped necrotic lesions on leaves."],
     ["Transfer learning",
      "Re-using a network pretrained on a large source corpus (ImageNet) as a feature extractor or "
      "initialisation for a smaller target task."],
     ["On-device inference",
      "Running model prediction locally on the user's hardware (a smartphone CPU) without a server "
      "or internet connection."],
     ["Leakage-free (group-aware) split",
      "A train/test partition in which all images derived from the same source leaf are kept "
      "together, preventing inflated scores from near-duplicate leakage."],
     ["Grad-CAM",
      "Gradient-weighted Class Activation Mapping; a method that highlights the image regions most "
      "responsible for a prediction (Tjoa and Guan, 2021)."]])
p("Interpretability is essential before automated diagnoses influence the decisions of farmers. "
  "Grad-CAM (Tjoa and Guan, 2021) localises the evidence behind a prediction, allowing "
  "verification that a blight call rests on lesion tissue rather than background artefacts, the "
  "agricultural analogue of the auditor-facing transparency required in other regulated AI domains. "
  "The literature therefore motivates the three pillars of the present design: transfer learning "
  "with a mobile backbone, leakage-free evaluation that anticipates the generalisation gap, and "
  "attribution-based interpretability.")

# ---------------- 3 Methodology
h1("4 Methodology and instruments")
p("A prototype-demonstration approach was adopted using field-collected maize-leaf photographs "
  "representative of conditions a deployed application would encounter. The working dataset contains "
  "1,821 images, each centre-cropped to a square and resized to 224x224 pixels. The positive class "
  "is NCLB; the negative class consists of other maize diseases. Every image is tagged with the "
  "identifier of the source leaf it derives from, and this identifier is used as the grouping key "
  "for partitioning. Data-quality checks confirmed that every labelled file is present and that no "
  "source leaf spans more than one partition.")
tbl("Table 2. Dataset profile and partition.",
    ["Attribute", "Value", "Interpretation"],
    [["Images", "1,821", "Field-collected maize leaves (224x224x3)"],
     ["Class balance", "NCLB 827 (45.4%) | Other diseases 994 (54.6%)", "Near-balanced binary task"],
     ["Partition", "Train 1,299 | Val 273 | Test 249", "Group-aware, leakage-free"],
     ["Backbone", "MobileNetV3-Small (ImageNet)", "576-dimensional embedding"],
     ["Majority baseline", "54.6%", "Accuracy floor to beat"]])

h2("4.1 Data preparation and feature representation")
p("Each raw photograph is reduced to a fixed-size tensor by a deterministic preprocessing map and "
  "then represented in two complementary ways: by interpretable hand-crafted colour/texture "
  "descriptors, and by deep embeddings extracted from the pretrained backbone. The descriptors "
  "(green fraction, necrosis/tan fraction, saturation, intensity standard deviation, and Sobel edge "
  "density) provide a transparent baseline and exploratory signal; the embeddings supply the rich, "
  "general visual features that drive the classifiers.")
p("4.1.1 Preprocessing steps. Step 1: read the image and convert to RGB. Step 2: centre-crop to a "
  "square of side equal to the shorter image dimension. Step 3: resize to 224x224 with Lanczos "
  "resampling. Step 4: scale to [0,1] and standardise with ImageNet channel statistics. Step 5: for "
  "training only, apply augmentation (random horizontal/vertical flips, +/-20 degree rotation, and "
  "colour jitter), which is label-preserving because lesion appearance is orientation-invariant. "
  "Step 6: cache embeddings once per image for the classical route to avoid recomputation.")

h2("4.2 Modelling pipeline")
p("Route A (frozen embeddings). Every image is passed once through MobileNetV3-Small with its "
  "classifier removed, yielding a 576-dimensional embedding. Logistic regression, linear and RBF "
  "support-vector machines, a random forest, and histogram gradient boosting are trained on these "
  "embeddings (using the scikit-learn library); the best model is selected by validation ROC-AUC "
  "and evaluated once on the test set. Route B (fine-tuned CNN). The same backbone is fine-tuned "
  "end-to-end in PyTorch (Ansel et al., 2024): early blocks are frozen, the final blocks and a fresh "
  "two-class head are trained with class-weighted cross-entropy, the AdamW optimiser with cosine "
  "annealing, and early stopping on validation ROC-AUC. Both routes are evaluated on the identical "
  "group-aware partition.")

h2("4.3 Performance measures and quality control")
p("Model quality is assessed using accuracy, balanced accuracy, per-class and macro-averaged "
  "precision/recall/F1, ROC-AUC, PR-AUC, the Gini coefficient, and the confusion matrix (Hicks et "
  "al., 2022). Interpretability is assessed with Grad-CAM. A fixed random seed (42) governs "
  "splitting, augmentation and training for reproducibility, and the group-aware partition is "
  "asserted to contain no shared source leaf across splits.")
p("4.3.1 Assumptions. Representativeness: the field images are assumed broadly representative of "
  "the deployment setting, though a single collection style limits external validity. Label quality: "
  "class labels are treated as authoritative. Independence: the group-aware split assumes that "
  "leaves are independent units; spatial or temporal correlation in a live feed may warrant "
  "rolling-origin validation. Explainability: Grad-CAM attributions apply to the final convolutional "
  "block and are indicative rather than exhaustive.")
p("4.3.2 Reproducibility. Exploratory analysis and modelling are implemented in a single executable "
  "notebook; figures export to a numbered directory and headline metrics serialise to a JSON file. "
  "Replication requires Python 3.11+, PyTorch and torchvision, scikit-learn, and the listed "
  "dependencies, with the data partition fixed by the published label file.")

h2("4.4 Mathematical formulation and scoring equations")
p("This section states the diagnosis problem in explicit notation. Let N denote the number of "
  "images and C = 2 the number of classes. Each image i has raw form I_i and a binary label y_i = 1 "
  "if the leaf shows NCLB and 0 otherwise. The deterministic preprocessing map produces a "
  "normalised tensor; the backbone produces an embedding; a head produces class probabilities.")
tbl("Table 3. Mathematical notation.",
    ["Symbol", "Description"],
    [["N", "Number of images (working set N = 1,821)"],
     ["C", "Number of classes (C = 2: NCLB, other disease)"],
     ["I_i, x-tilde_i", "Raw image and its normalised 224x224x3 tensor"],
     ["z_i", "576-dim MobileNetV3 embedding of image i"],
     ["y_i", "Label: 1 if NCLB, else 0"],
     ["p_i", "Predicted class-probability vector (softmax output)"],
     ["w_c", "Inverse-frequency weight of class c"],
     ["theta", "Trainable network parameters"]])
p("Preprocessing and representation. With image width W_i and height H_i:")
eq(r"s_i = \min(W_i,\, H_i), \qquad \mathrm{crop}_i = \mathrm{CentreCrop}(I_i,\, s_i)", "(1)")
eq(r"x_i = \mathrm{Resize}_{224}(\mathrm{crop}_i)", "(2)")
eq(r"\tilde{x}_i = \frac{x_i/255 - \mu}{\sigma}, \qquad \mu=(0.485,\,0.456,\,0.406),\ \ \sigma=(0.229,\,0.224,\,0.225)", "(3)")
eq(r"z_i = g_{\theta}(\tilde{x}_i)\ \in\ \mathbb{R}^{576} \qquad (\mathrm{MobileNetV3\ backbone})", "(4)")
p("Classifier head and probabilities. The fine-tuned head is a small multilayer perceptron; for "
  "class c the softmax probability is:")
eq(r"o_i = W_2\,\mathrm{Hardswish}(W_1 z_i + b_1) + b_2", "(5)")
eq(r"p_{i,c} = \frac{\exp(o_{i,c})}{\sum_{c'=1}^{C}\exp(o_{i,c'})}", "(6)")
p("Objective. Training minimises class-weighted cross-entropy, which counters the mild class "
  "imbalance by up-weighting the rarer class:")
eq(r"\mathcal{L}(\theta) = -\frac{1}{N}\sum_{i=1}^{N}\sum_{c=1}^{C} w_c\,\mathbf{1}[\,y_i = c\,]\,\log p_{i,c}", "(7)")
eq(r"w_c = \frac{N}{C\,N_c}", "(8)")
p("On the training partition (N = 1,299; NCLB N_1 = 577, other N_0 = 722) Equation (8) gives "
  "w_NCLB = 1,299 / (2 x 577) = 1.126 and w_other = 1,299 / (2 x 722) = 0.900, so blight examples "
  "carry about 25% more loss weight than other-disease examples.")
p("Discrimination and calibration. From the confusion-matrix counts (true/false positives and "
  "negatives) the operating metrics are:")
eq(r"\mathrm{Accuracy} = \frac{TP+TN}{TP+TN+FP+FN}", "(9)")
eq(r"\mathrm{Precision} = \frac{TP}{TP+FP}, \qquad \mathrm{Recall} = \frac{TP}{TP+FN}", "(10)")
eq(r"F_1 = \frac{2\,\mathrm{Precision}\cdot\mathrm{Recall}}{\mathrm{Precision}+\mathrm{Recall}}", "(11)")
eq(r"F_1^{\mathrm{macro}} = \frac{1}{C}\sum_{c=1}^{C} F_{1,c}", "(12)")
eq(r"\mathrm{BalancedAcc} = \frac{1}{2}\!\left(\frac{TP}{TP+FN} + \frac{TN}{TN+FP}\right)", "(13)")
eq(r"\mathrm{TPR}(t) = \frac{TP}{TP+FN}, \qquad \mathrm{FPR}(t) = \frac{FP}{FP+TN}", "(14)")
eq(r"\mathrm{AUROC} = \int_0^1 \mathrm{TPR}\!\left(\mathrm{FPR}^{-1}(u)\right)du, \qquad \mathrm{Gini} = 2\,\mathrm{AUROC} - 1", "(15)")
p("Interpretability. Grad-CAM weights the final convolutional feature maps A_k by the gradient of "
  "the NCLB logit y^c, producing a localisation heat-map (Tjoa and Guan, 2021):")
eq(r"\alpha_k^{c} = \frac{1}{Z}\sum_u\sum_v \frac{\partial y^{c}}{\partial A_k(u,v)}", "(16)")
eq(r"L_{\mathrm{GradCAM}}^{c} = \mathrm{ReLU}\!\left(\sum_k \alpha_k^{c}\, A_k\right)", "(17)")
p("Worked example. The fine-tuned model's test confusion matrix is TP = 111, FN = 8, FP = 2, "
  "TN = 128 over 249 images. Substituting into Equations (9)-(13): Accuracy = (111 + 128) / 249 = "
  "0.9598; NCLB Recall = 111 / 119 = 0.9328; NCLB Precision = 111 / 113 = 0.9823; F1 = 2 x 0.9823 x "
  "0.9328 / (0.9823 + 0.9328) = 0.9570; Balanced Accuracy = (0.9328 + 128/130) / 2 = 0.9587. With "
  "ROC-AUC = 0.989, Equation (15) gives Gini = 2 x 0.989 - 1 = 0.978.")
tbl("Table 4. Calculated evaluation metrics (held-out test set).",
    ["Metric", "Equation", "Value"],
    [["Accuracy", "(9)", "0.9598"],
     ["Balanced accuracy", "(13)", "0.9587"],
     ["Macro-F1", "(12)", "0.9597"],
     ["NCLB precision", "(10)", "0.9823"],
     ["NCLB recall", "(10)", "0.9328"],
     ["NCLB F1", "(11)", "0.9570"],
     ["ROC-AUC", "(15)", "0.989"],
     ["PR-AUC", "area under P-R curve", "0.987"],
     ["Gini coefficient", "(15)", "0.978"]])

# ---------------- 4 Results
h1("5 Exploratory analysis and prototype results")
p("Exploratory analysis confirmed that NCLB lesions form a distinctive tan, necrotic signature, but "
  "that on the harder NCLB-versus-other-disease task no single hand-crafted descriptor separates "
  "the classes well. Table 5 reports the univariate ROC-AUC of each descriptor (direction-agnostic). "
  "The strongest, mean green intensity, reaches only 0.727, and necrosis fraction 0.633, because "
  "other diseases also produce discolouration. This motivates deep features: the 576-dimensional "
  "embeddings separate the classes almost completely, as the projection in Figure 3 shows.")
tbl("Table 5. Univariate ROC-AUC of hand-crafted descriptors for NCLB.",
    ["Descriptor", "ROC-AUC", "Interpretation"],
    [["Mean green intensity", "0.727", "Weak-moderate"],
     ["Edge density (Sobel)", "0.684", "Weak"],
     ["Green fraction", "0.662", "Weak"],
     ["Mean red intensity", "0.637", "Weak"],
     ["Necrosis (tan) fraction", "0.633", "Weak"],
     ["Random baseline", "0.500", "Chance"]])
fig("outputs/figures/05_embedding_projection.png",
    "Figure 2. Two-dimensional projections (t-SNE and PCA) of the 576-dimensional MobileNetV3 "
    "embeddings, coloured by class. The classes occupy largely distinct regions before any "
    "fine-tuning. Source: study notebook.")

h2("5.1 Model performance")
p("On frozen embeddings, every classical model exceeded ROC-AUC 0.94 in validation; the RBF "
  "support-vector machine was selected (validation ROC-AUC 0.987) and achieved test ROC-AUC 0.998, "
  "confirming that ImageNet features alone separate the classes strongly. The fine-tuned CNN, the "
  "intended on-device model, achieved test accuracy 96.0%, balanced accuracy 0.959, macro-F1 0.960 "
  "and ROC-AUC 0.989, early-stopping after 12 epochs. Table 6 compares the models and Figures 4-6 "
  "report the training dynamics and test discrimination.")
tbl("Table 6. Model comparison.",
    ["Model", "ROC-AUC", "Role"],
    [["RBF-SVM on frozen embeddings", "0.998 (test)", "Strong baseline"],
     ["Histogram gradient boosting", "0.968 (val)", "Embedding baseline"],
     ["Logistic regression", "0.962 (val)", "Embedding baseline"],
     ["Linear SVM", "0.960 (val)", "Embedding baseline"],
     ["Random forest", "0.942 (val)", "Embedding baseline"],
     ["Fine-tuned MobileNetV3 (CNN)", "0.989 (test)", "Production on-device model"],
     ["Majority-class baseline", "0.500", "Reference"]])
fig("outputs/figures/06_classical_auc.png",
    "Figure 3. Validation ROC-AUC of classical classifiers trained on frozen embeddings. "
    "Source: study notebook.")
fig("outputs/figures/07_training_curves.png",
    "Figure 4. Fine-tuned CNN training loss and validation ROC-AUC/accuracy across epochs. "
    "Source: study notebook.")
fig("outputs/figures/08_test_metrics_paper.png",
    "Figure 5. Held-out test confusion matrix of the fine-tuned model (accuracy 0.960). "
    "Source: study notebook.")
p("A per-collection breakdown (Table 7, Figure 7) localises the residual error: other-disease "
  "leaves are almost never mistaken for blight (98.5% correct), whereas the harder side is NCLB "
  "itself, where 8 of 119 cases are missed (93.3% correct). The model's bias is therefore towards "
  "missing faint or early blight rather than over-calling it, an operating point that a deployed "
  "threshold can be tuned to adjust according to the agronomic cost of a missed infection.")
tbl("Table 7. Test accuracy by source collection.",
    ["Source collection", "Test images", "Accuracy"],
    [["Other diseases", "130", "98.5%"],
     ["NCLB", "119", "93.3%"]])
fig("outputs/figures/11_per_dataset.png",
    "Figure 6. Test accuracy by source collection. Source: study notebook.")

h2("5.2 Interpretability and offline deployment")
p("Grad-CAM attributions on confident test predictions concentrate on lesion tissue, supporting "
  "the claim that the network keys on disease morphology rather than background; the attributions "
  "are somewhat diffuse, which is consistent with the modest dataset and motivates a future "
  "independently-collected validation set. The intended deployment is fully offline. The trained "
  "weights are quantised to 8-bit integers and converted to a mobile format of only a few "
  "megabytes, bundled into the application, and executed on the phone CPU. Figures 8-10 present the "
  "system architecture, the on-device inference-and-recommendation pipeline, and the train-to-"
  "deploy lifecycle.")
fig("docs/architecture/01_system_architecture.png",
    "Figure 7. Layered offline system architecture: all core functions run on the device; cloud "
    "connectivity is optional. Source: study design.")
fig("docs/architecture/02_inference_pipeline.png",
    "Figure 8. On-device inference and recommendation pipeline, including a low-confidence "
    "deferral to a human extension officer. Source: study design.")
fig("docs/architecture/03_model_lifecycle.png",
    "Figure 9. Model lifecycle: one-off training and quantised export, then repeated offline "
    "on-device inference. Source: study design.")

# ---------------- 5 Discussions
h1("6 Discussions")
p("The results establish that distinguishing NCLB from other maize diseases, a substantially harder "
  "task than separating diseased from healthy leaves, is feasible at 96% accuracy with a model "
  "small enough to run offline on a phone. The error structure is informative: discrimination is "
  "near-perfect for the other-disease class and the residual difficulty lies in recall of blight, "
  "where faint or early lesions are occasionally missed. For a screening tool whose purpose is to "
  "prompt timely action, the decision threshold can be lowered to raise blight recall at the cost "
  "of more false alarms, a trade-off that should be calibrated to local agronomic economics.")
p("The headline in-dataset metrics are strong but must be interpreted with discipline. As the "
  "plant-disease literature repeatedly shows (Ngugi et al., 2021; Shoaib et al., 2023), models that "
  "excel on a single image collection can degrade markedly on imagery from a new field, season, or "
  "camera. The leakage-free, group-aware partition used here rules out near-duplicate inflation, "
  "but it cannot substitute for an independently photographed validation set, which is the most "
  "important next step before field deployment. The Grad-CAM evidence that predictions rest on "
  "lesions is reassuring but not conclusive given the diffuse attributions observed.")
p("Practical and policy implications are immediate. A compact, quantised model delivered inside an "
  "offline application removes the connectivity barrier that excludes rural smallholders from "
  "existing tools, and multilingual delivery in English, Shona and Ndebele addresses literacy and "
  "language diversity. A phased rollout is recommended. Phase 1 (0-6 months): package the model in "
  "a pilot application, validate against agronomist labels in two districts, and instrument a "
  "low-confidence deferral to extension officers. Phase 2 (6-18 months): expand to additional maize "
  "diseases as a multi-class model, add an independently collected cross-field test set, and gather "
  "farmer feedback. Phase 3 (18-36 months): integrate a treatment knowledge base with localised "
  "agrochemical and natural-remedy guidance, and establish a governed, opt-in channel for model "
  "improvement when connectivity is available.")
p("Ethical deployment requires that the tool advises rather than dictates: low-confidence cases are "
  "deferred to humans, version history is logged, and the system never substitutes for "
  "agronomic judgement. These safeguards mirror responsible-AI guidance for public-facing systems "
  "and are consistent with Zimbabwe's National AI Strategy (2026-2030).")

# ---------------- 6 Conclusions
h1("7 Conclusions")
p("This paper demonstrated an offline, on-device artificial-intelligence pipeline for precision "
  "crop disease diagnosis, specialised to Northern Corn Leaf Blight in maize and evaluated on the "
  "demanding task of separating NCLB from other maize diseases. Transfer learning on a lightweight "
  "MobileNetV3-Small backbone yielded a fine-tuned model with test accuracy 96.0%, balanced "
  "accuracy 0.959 and ROC-AUC 0.989, while a support-vector classifier on frozen embeddings reached "
  "test ROC-AUC 0.998; Grad-CAM confirmed lesion-focused predictions, and a leakage-free group-aware "
  "protocol ensured the scores reflect generalisation to unseen leaves. Because the quantised model "
  "is a few megabytes and runs on a CPU, the diagnostic engine is deployable entirely offline on "
  "low-cost smartphones, with multilingual treatment guidance.")
p("Future work centres on external validity: an independently collected cross-field test set, "
  "extension to a multi-class model spanning the principal maize diseases and their look-alikes, "
  "field trials measuring the effect on time-to-treatment and yield, and a governed continuous-"
  "learning channel. By embedding 4IR technologies in an inclusive, offline-first design, the work "
  "shows how precision agriculture can reduce crop losses and strengthen food security in "
  "resource-constrained environments, advancing SDG 2 and SDG 9 and Zimbabwe's National AI "
  "Strategy (2026-2030).")

# ---------------- references
refs([
 "Alzubaidi, L., Zhang, J., Humaidi, A.J., Al-Dujaili, A., Duan, Y., Al-Shamma, O., Santamaria, J., Fadhel, M.A., Al-Amidie, M. and Farhan, L. (2021) 'Review of deep learning: concepts, CNN architectures, challenges, applications, future directions', Journal of Big Data, 8, 53. doi: 10.1186/s40537-021-00444-8.",
 "Ansel, J. et al. (2024) 'PyTorch 2: Faster machine learning through dynamic Python bytecode transformation and graph compilation', Proceedings of ACM ASPLOS, pp. 929-947. doi: 10.1145/3620665.3640366.",
 "CABI (2024) Fall armyworm. Centre for Agriculture and Biosciences International. Available at: https://www.cabi.org/what-we-do/fall-armyworm/ (Accessed: 27 June 2026).",
 "CIPESA (2026) Zimbabwe's National AI Strategy: Policy Lessons for Africa. Collaboration on International ICT Policy for East and Southern Africa. Available at: https://cipesa.org/2026/05/zimbabwes-national-ai-strategy-policy-lessons-for-africa/ (Accessed: 27 June 2026).",
 "Day, R., Abrahams, P., Bateman, M., Beale, T., Clottey, V., Cock, M., Colmenarez, Y., Corniani, N., Early, R., Godwin, J., Gomez, J., Moreno, P.G., Murphy, S.T., Oppong-Mensah, B., Phiri, N., Pratt, C., Silvestri, S. and Witt, A. (2017) 'Fall armyworm: impacts and implications for Africa', Outlooks on Pest Management, 28(5), pp. 196-201. doi: 10.1564/v28_oct_02.",
 "Dube, Z.L., Moyo, F., Mjimba, V., Murai, T. and Ndlovu, M.E. (2025) 'Bridging the readiness gap: Zimbabwe's AI capacity for economic development', African Journal of Advances in Engineering and Technology, 1(2). doi: 10.31920/2978-3224/2025/v1n2a4.",
 "FAO (2021a) The State of Food and Agriculture 2021. Making agrifood systems more resilient to shocks and stresses. Rome: Food and Agriculture Organization. doi: 10.4060/cb4476en.",
 "FAO (2021b) Plant Production and Protection: up to 40 percent of food crops lost to pests and diseases annually. Food and Agriculture Organization. Available at: https://www.fao.org/plant-production-protection/about/en (Accessed: 27 June 2026).",
 "Government of Zimbabwe (2026) National Artificial Intelligence Strategy 2026-2030. Harare: Ministry of Information Communication Technology, Postal and Courier Services. Available at: https://veritaszim.net/sites/veritas_d/files/Zimbabwe%20National%20Artificial%20Intelligence%20Strategy.pdf (Accessed: 27 June 2026).",
 "Hicks, S.A., Strumke, I., Thambawita, V., Hammou, M., Riegler, M.A., Halvorsen, P. and Parasa, S. (2022) 'On evaluation metrics for medical applications of artificial intelligence', Scientific Reports, 12, 5979. doi: 10.1038/s41598-022-09954-8.",
 "Jackson-Ziems, T.A. (2014) Northern Corn Leaf Blight. University of Nebraska-Lincoln Extension. Available at: https://digitalcommons.unl.edu/plantpathpapers/550/ (Accessed: 27 June 2026).",
 "Mensah, P.K., Akoto-Adjepong, V., Adu, K., Ayidzoe, M.A., Bediako, E.A., Nyarko-Boateng, O., Boateng, S. and Donkor, E.F. (2023) 'CCMT: Dataset for crop pest and disease detection', Data in Brief, 49, 109306. doi: 10.1016/j.dib.2023.109306.",
 "Mhlanga, D. (2021) 'Artificial Intelligence in the Industry 4.0, and Its Impact on Poverty, Innovation, Infrastructure Development, and the Sustainable Development Goals: Lessons from Emerging Economies?', Sustainability, 13(11), 5788. doi: 10.3390/su13115788.",
 "Milling MEA (2024) 'Country focus: Zimbabwe's agricultural crisis in 2023/24', Milling Middle East & Africa. Available at: https://millingmea.com/ (Accessed: 27 June 2026).",
 "Ngugi, L.C., Abelwahab, M. and Abo-Zahhad, M. (2021) 'Recent advances in image processing techniques for automated leaf pest and disease recognition - A review', Information Processing in Agriculture, 8(1), pp. 27-51. doi: 10.1016/j.inpa.2020.04.004.",
 "OCHA (2024) Zimbabwe Drought Flash Appeal, May 2024 - April 2025. United Nations Office for the Coordination of Humanitarian Affairs. Available at: https://www.unocha.org/ (Accessed: 27 June 2026).",
 "Rwomushana, I., Bateman, M., Beale, T., Beseh, P., Cameron, K., Chiluba, M., Clottey, V., Davis, T., Day, R., Early, R., Godwin, J., Gonzalez-Moreno, P., Kansiime, M., Kenis, M., Makale, F., Mugambi, I., Murphy, S., Nunda, W., Phiri, N., Pratt, C. and Tambo, J. (2018) Fall armyworm: impacts and implications for Africa - Evidence Note Update, October 2018. Wallingford: CABI.",
 "Sambasivam, G. and Opiyo, G.D. (2021) 'A predictive machine learning application in agriculture: Cassava disease detection and classification with imbalanced dataset using convolutional neural networks', Egyptian Informatics Journal, 22(1), pp. 27-34. doi: 10.1016/j.eij.2020.02.007.",
 "Savary, S., Willocquet, L., Pethybridge, S.J., Esker, P., McRoberts, N. and Nelson, A. (2019) 'The global burden of pathogens and pests on major food crops', Nature Ecology & Evolution, 3(3), pp. 430-439. doi: 10.1038/s41559-018-0793-y.",
 "Shoaib, M., Shah, B., Ei-Sappagh, S., Ali, A., Ullah, A., Alenezi, F., Gechev, T., Hussain, T. and Ali, F. (2023) 'An advanced deep learning models-based plant disease detection: A review of recent research', Frontiers in Plant Science, 14, 1158933. doi: 10.3389/fpls.2023.1158933.",
 "Tambo, J.A., Kansiime, M.K., Mugambi, I., Rwomushana, I., Kenis, M., Day, R.K. and Lamontagne-Godwin, J. (2020) 'Understanding smallholders' responses to fall armyworm (Spodoptera frugiperda) invasion: evidence from five African countries', Food and Energy Security, 9(3), e224. doi: 10.1002/fes3.224.",
 "Tan, M. and Le, Q.V. (2021) 'EfficientNetV2: Smaller models and faster training', Proceedings of ICML, pp. 10096-10106.",
 "Tjoa, E. and Guan, C. (2021) 'A survey on explainable artificial intelligence (XAI): toward medical XAI', IEEE Transactions on Neural Networks and Learning Systems, 32(11), pp. 4793-4813. doi: 10.1109/TNNLS.2020.3027314.",
 "USDA FAS (2024) Zimbabwe Grain and Feed Annual. United States Department of Agriculture, Foreign Agricultural Service. Available at: https://fas.usda.gov/ (Accessed: 27 June 2026).",
 "Waldamichael, F.G., Debelee, T.G., Schwenker, F., Ayano, Y.M. and Kebede, S.R. (2022) 'Machine Learning in Cereal Crops Disease Detection: A Review', Algorithms, 15(3), 75. doi: 10.3390/a15030075.",
 "Zhuang, F., Qi, Z., Duan, K., Xi, D., Zhu, Y., Zhu, H., Xiong, H. and He, Q. (2021) 'A comprehensive survey on transfer learning', Proceedings of the IEEE, 109(1), pp. 43-76. doi: 10.1109/JPROC.2020.3004555.",
 "Crop Disease Detector Project (2026) Maize NCLB image dataset and offline diagnosis pipeline. Available at: https://github.com/ (Accessed: 26 June 2026).",
])
B.append(("h1", "8 Acknowledgements"))
p("The authors thank the agricultural extension officers and smallholder farmers whose field "
  "knowledge informed dataset curation and problem framing, and colleagues in data science whose "
  "review improved the evaluation protocol. This prototype demonstration was developed under the "
  "theme Empowering SMEs with 4IR and AI for Inclusive and Sustainable Growth. Responsibility for "
  "any errors and omissions remains with the authors.")

# ============================================================================= equation typesetting
def prerender_equations():
    """Render each LaTeX equation to a transparent PNG via matplotlib mathtext."""
    import matplotlib; matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    eqdir="outputs/figures/eq"; os.makedirs(eqdir, exist_ok=True)
    idx=0
    for n, blk in enumerate(B):
        if blk[0]=="eq":
            idx+=1; latex, num = blk[1], blk[2]
            path=f"{eqdir}/eq{idx:02d}.png"
            fig=plt.figure(figsize=(0.01,0.01))
            fig.text(0,0, f"${latex}$", fontsize=11, color="black")
            fig.savefig(path, dpi=300, transparent=True, bbox_inches="tight", pad_inches=0.03)
            plt.close(fig)
            w,h=PILImage.open(path).size
            B[n]=("eq", latex, num, path, w/300.0, h/300.0)
    print("rendered", idx, "equations")

# ============================================================================= DOCX renderer
def render_docx():
    import docx
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin=Inches(1); s.bottom_margin=Inches(1)
        s.left_margin=Inches(1); s.right_margin=Inches(1)
    base = doc.styles["Normal"]; base.font.name="Times New Roman"; base.font.size=Pt(11)
    base.element.rPr.rFonts.set(qn("w:eastAsia"),"Times New Roman")

    def para(text, align=AL.JUSTIFY, size=11, bold=False, italic=False, before=2, after=6, color=None):
        pr=doc.add_paragraph(); pr.alignment=align
        pr.paragraph_format.space_before=Pt(before); pr.paragraph_format.space_after=Pt(after)
        r=pr.add_run(text); r.bold=bold; r.italic=italic; r.font.size=Pt(size)
        if color: r.font.color.rgb=RGBColor(*color)
        return pr

    for blk in B:
        k=blk[0]
        if k=="hdr":
            for i,l in enumerate(blk[1]): para(l, AL.CENTER, 9.5, bold=(i==0), after=1, before=1)
        elif k=="title":
            para(blk[1], AL.CENTER, 15, bold=True, before=10, after=8)
        elif k=="auth":
            for i,l in enumerate(blk[1]): para(l, AL.CENTER, (11 if i==0 else 9.5), after=1, before=1)
        elif k=="meta":
            for l in blk[1]: para(l, AL.CENTER, 9.5, italic=True, after=1, before=1)
        elif k=="abstract":
            pr=doc.add_paragraph(); pr.alignment=AL.JUSTIFY
            pr.paragraph_format.space_before=Pt(10)
            r=pr.add_run("Abstract. "); r.bold=True; r.font.size=Pt(10.5)
            r2=pr.add_run(blk[1]); r2.font.size=Pt(10.5)
        elif k=="keywords":
            pr=doc.add_paragraph(); pr.alignment=AL.JUSTIFY
            r=pr.add_run("Key words: "); r.bold=True; r.font.size=Pt(10.5)
            r2=pr.add_run(blk[1]); r2.italic=True; r2.font.size=Pt(10.5)
        elif k=="h1":
            para(blk[1], AL.LEFT, 12.5, bold=True, before=12, after=4)
        elif k=="h2":
            para(blk[1], AL.LEFT, 11, bold=True, italic=True, before=8, after=3)
        elif k=="body":
            para(blk[1])
        elif k=="eq":
            num = blk[2]; imgpath = blk[3] if len(blk)>5 else None; win = blk[4] if len(blk)>5 else None
            pr=doc.add_paragraph(); pr.alignment=AL.LEFT
            pr.paragraph_format.space_before=Pt(4); pr.paragraph_format.space_after=Pt(4)
            pr.paragraph_format.left_indent=Inches(0.4)
            from docx.enum.text import WD_TAB_ALIGNMENT
            pr.paragraph_format.tab_stops.add_tab_stop(Inches(6.0), WD_TAB_ALIGNMENT.RIGHT)
            if imgpath and os.path.exists(imgpath):
                pr.add_run().add_picture(imgpath, width=Inches(min(win, 5.2)))
            else:
                r=pr.add_run(blk[1]); r.italic=True; r.font.size=Pt(10.5)
            r2=pr.add_run("\t"+num); r2.font.size=Pt(10.5)
        elif k=="table":
            cap, head, rows = blk[1], blk[2], blk[3]
            para(cap, AL.LEFT, 10, bold=True, before=8, after=2)
            t=doc.add_table(rows=1, cols=len(head)); t.style="Light Grid Accent 1"
            t.alignment=WD_TABLE_ALIGNMENT.CENTER
            for j,h in enumerate(head):
                c=t.rows[0].cells[j]; c.paragraphs[0].runs and None
                rr=c.paragraphs[0].add_run(h); rr.bold=True; rr.font.size=Pt(9.5)
            for row in rows:
                cells=t.add_row().cells
                for j,val in enumerate(row):
                    rr=cells[j].paragraphs[0].add_run(str(val)); rr.font.size=Pt(9.5)
            doc.add_paragraph().paragraph_format.space_after=Pt(2)
        elif k=="figure":
            path, cap = blk[1], blk[2]
            if os.path.exists(path):
                pr=doc.add_paragraph(); pr.alignment=AL.CENTER
                pr.paragraph_format.space_before=Pt(6)
                w,h=PILImage.open(path).size; iw=min(6.0, 6.0);
                pr.add_run().add_picture(path, width=Inches(6.0 if w>=h else 3.6))
            para(cap, AL.LEFT, 9, italic=True, before=2, after=8)
        elif k=="refs":
            para("References", AL.LEFT, 12.5, bold=True, before=12, after=4)
            for i,r in enumerate(blk[1],1):
                pr=doc.add_paragraph(); pr.alignment=AL.JUSTIFY
                pr.paragraph_format.left_indent=Inches(0.3); pr.paragraph_format.first_line_indent=Inches(-0.3)
                pr.paragraph_format.space_after=Pt(3)
                rr=pr.add_run(f"[{i}] {r}"); rr.font.size=Pt(9.5)
    doc.save(DOCX_OUT)
    print("wrote", DOCX_OUT)

# ============================================================================= PDF renderer
def render_pdf():
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import inch
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
                                    Image as RLImage, KeepTogether)
    import xml.sax.saxutils as sx
    def esc(t): return sx.escape(str(t))
    styles=getSampleStyleSheet()
    body=ParagraphStyle("body", parent=styles["Normal"], fontName="Times-Roman", fontSize=10.5,
                        leading=14.5, alignment=TA_JUSTIFY, spaceAfter=6)
    H=lambda **k: ParagraphStyle(**{"parent":body, **k})
    st_hdr=H(name="h", alignment=TA_CENTER, fontSize=9.5, leading=12, spaceAfter=1)
    st_hdrb=H(name="hb", alignment=TA_CENTER, fontSize=9.5, leading=12, fontName="Times-Bold", spaceAfter=1)
    st_title=H(name="t", alignment=TA_CENTER, fontName="Times-Bold", fontSize=15, leading=18, spaceBefore=10, spaceAfter=8)
    st_auth=H(name="a", alignment=TA_CENTER, fontSize=10.5, leading=13, spaceAfter=1)
    st_meta=H(name="m", alignment=TA_CENTER, fontSize=9.5, leading=12, fontName="Times-Italic", spaceAfter=1)
    st_h1=H(name="H1", fontName="Times-Bold", fontSize=12.5, leading=15, spaceBefore=12, spaceAfter=4, alignment=TA_LEFT)
    st_h2=H(name="H2", fontName="Times-BoldItalic", fontSize=11, leading=14, spaceBefore=8, spaceAfter=3, alignment=TA_LEFT)
    st_eq=H(name="eq", fontName="Times-Italic", fontSize=10.5, leading=14, leftIndent=28, alignment=TA_LEFT, spaceBefore=2, spaceAfter=2)
    st_cap=H(name="cap", fontName="Times-Italic", fontSize=9, leading=12, spaceBefore=2, spaceAfter=8, alignment=TA_LEFT)
    st_ref=H(name="ref", fontSize=9.5, leading=12, leftIndent=18, firstLineIndent=-18, spaceAfter=3, alignment=TA_JUSTIFY)

    flow=[]; usable=A4[0]-2*0.9*inch
    for blk in B:
        k=blk[0]
        if k=="hdr":
            for i,l in enumerate(blk[1]): flow.append(Paragraph(esc(l), st_hdrb if i==0 else st_hdr))
        elif k=="title": flow.append(Paragraph(esc(blk[1]), st_title))
        elif k=="auth":
            for i,l in enumerate(blk[1]): flow.append(Paragraph(esc(l), st_auth))
        elif k=="meta":
            for l in blk[1]: flow.append(Paragraph(esc(l), st_meta))
        elif k=="abstract":
            flow.append(Spacer(1,6)); flow.append(Paragraph("<b>Abstract.</b> "+esc(blk[1]),
                        H(name="ab", parent=body, fontSize=10.5)))
        elif k=="keywords":
            flow.append(Paragraph("<b>Key words:</b> <i>"+esc(blk[1])+"</i>", H(name="kw", parent=body, fontSize=10.5)))
        elif k=="h1": flow.append(Paragraph(esc(blk[1]), st_h1))
        elif k=="h2": flow.append(Paragraph(esc(blk[1]), st_h2))
        elif k=="body": flow.append(Paragraph(esc(blk[1]), body))
        elif k=="eq":
            num=blk[2]; imgpath=blk[3] if len(blk)>5 else None
            win=blk[4] if len(blk)>5 else None; hin=blk[5] if len(blk)>5 else None
            if imgpath and os.path.exists(imgpath):
                maxw=usable-0.55*inch; w_in=win; h_in=hin
                if w_in*inch>maxw:
                    sc=maxw/(w_in*inch); w_in*=sc; h_in*=sc
                numpar=Paragraph("<b>"+esc(num)+"</b>", H(name="eqn", fontSize=10.5, alignment=TA_RIGHT))
                t=Table([[RLImage(imgpath, width=w_in*inch, height=h_in*inch), numpar]],
                        colWidths=[usable-0.55*inch, 0.55*inch])
                t.setStyle(TableStyle([("VALIGN",(0,0),(-1,-1),"MIDDLE"),
                    ("LEFTPADDING",(0,0),(0,0),26),("LEFTPADDING",(1,0),(1,0),0),
                    ("RIGHTPADDING",(0,0),(-1,-1),0),("TOPPADDING",(0,0),(-1,-1),3),
                    ("BOTTOMPADDING",(0,0),(-1,-1),3)]))
                flow.append(t)
            else:
                flow.append(Paragraph(esc(blk[1])+"&nbsp;&nbsp;&nbsp;&nbsp;<b>"+esc(num)+"</b>", st_eq))
        elif k=="table":
            cap, head, rows = blk[1], blk[2], blk[3]
            flow.append(Paragraph("<b>"+esc(cap)+"</b>", H(name="tc", parent=body, fontSize=10, spaceBefore=8, spaceAfter=3)))
            ncol=len(head)
            cellstyle=H(name="cell", parent=body, fontSize=8.8, leading=11, alignment=TA_LEFT, spaceAfter=0)
            headstyle=H(name="hcell", parent=cellstyle, fontName="Times-Bold")
            data=[[Paragraph(esc(h), headstyle) for h in head]]
            for row in rows: data.append([Paragraph(esc(v), cellstyle) for v in row])
            # column widths: first col wider for 2-col definition tables
            if ncol==2: cw=[usable*0.32, usable*0.68]
            elif ncol==3: cw=[usable*0.34, usable*0.18, usable*0.48]
            else: cw=[usable/ncol]*ncol
            t=Table(data, colWidths=cw, repeatRows=1)
            t.setStyle(TableStyle([
                ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#3D7EA6")),
                ("TEXTCOLOR",(0,0),(-1,0),colors.white),
                ("GRID",(0,0),(-1,-1),0.5,colors.HexColor("#9bb7c9")),
                ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white, colors.HexColor("#eef4f8")]),
                ("VALIGN",(0,0),(-1,-1),"TOP"),
                ("LEFTPADDING",(0,0),(-1,-1),5),("RIGHTPADDING",(0,0),(-1,-1),5),
                ("TOPPADDING",(0,0),(-1,-1),3),("BOTTOMPADDING",(0,0),(-1,-1),3),
            ]))
            flow.append(t); flow.append(Spacer(1,8))
        elif k=="figure":
            path, cap = blk[1], blk[2]
            if os.path.exists(path):
                w,h=PILImage.open(path).size
                tw=usable*(1.0 if w>=h else 0.55); th=tw*h/w
                maxh=4.2*inch
                if th>maxh: th=maxh; tw=th*w/h
                flow.append(Spacer(1,4))
                flow.append(KeepTogether([RLImage(path, width=tw, height=th),
                                          Paragraph(esc(cap), st_cap)]))
        elif k=="refs":
            flow.append(Paragraph("References", st_h1))
            for i,r in enumerate(blk[1],1):
                flow.append(Paragraph(f"[{i}] "+esc(r), st_ref))
    doc=SimpleDocTemplate(PDF_OUT, pagesize=A4, topMargin=0.9*inch, bottomMargin=0.9*inch,
                          leftMargin=0.9*inch, rightMargin=0.9*inch, title=TITLE)
    doc.build(flow)
    print("wrote", PDF_OUT)

def renumber_figures():
    """Rewrite the 'Figure N.' prefix in each figure caption to its running order."""
    import re
    k=0
    for n, blk in enumerate(B):
        if blk[0]=="figure":
            k+=1
            B[n]=("figure", blk[1], re.sub(r"^Figure\s+\d+\.", f"Figure {k}.", blk[2]))

if __name__=="__main__":
    renumber_figures()
    prerender_equations()
    render_docx()
    render_pdf()
